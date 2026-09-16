"""
Compiles all crawled + summarized page data into a single, organized
Word document (.docx), with optional embedded screenshots.
"""
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import Config


def build_document(pages: list[dict]):
    doc = Document()

    # Title page
    title = doc.add_heading("Site Feature Documentation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Base URL: {Config.BASE_URL}\n").italic = True
    meta.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n").italic = True
    meta.add_run(f"Pages documented: {len(pages)}").italic = True
    doc.add_page_break()

    # Table of contents (manual, since python-docx has no native TOC field rendering)
    doc.add_heading("Table of Contents", level=1)
    for i, p in enumerate(pages, 1):
        doc.add_paragraph(f"{i}. {p['title'] or p['url']}", style="List Number")
    doc.add_page_break()

    # One section per page
    for i, p in enumerate(pages, 1):
        doc.add_heading(f"{i}. {p['title'] or 'Untitled Page'}", level=1)

        url_para = doc.add_paragraph()
        url_run = url_para.add_run(f"URL: {p['url']}")
        url_run.font.size = Pt(9)
        url_run.italic = True

        # LLM/fallback summary — rendered as simple markdown-ish text
        _add_markdown_ish(doc, p["summary"])

        # Screenshot
        if p.get("screenshot_path") and os.path.exists(p["screenshot_path"]):
            try:
                doc.add_picture(p["screenshot_path"], width=Inches(6))
            except Exception:
                pass

        doc.add_page_break()

    os.makedirs(os.path.dirname(Config.OUTPUT_DOCX_PATH) or ".", exist_ok=True)
    doc.save(Config.OUTPUT_DOCX_PATH)
    print(f"[doc_generator] Saved documentation to {Config.OUTPUT_DOCX_PATH}")
    return Config.OUTPUT_DOCX_PATH


def _add_markdown_ish(doc, text: str):
    """Very small markdown->docx renderer: handles ###, **bold**, and - bullets."""
    if not text:
        doc.add_paragraph("(No content extracted for this page.)")
        return

    for line in text.split("\n"):
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            p = doc.add_paragraph()
            _add_bold_runs(p, line)


def _add_bold_runs(paragraph, line: str):
    """Splits on **bold** markers and adds runs accordingly."""
    parts = line.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if idx % 2 == 1:
            run.bold = True